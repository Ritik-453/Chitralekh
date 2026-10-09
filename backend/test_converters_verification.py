import io
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfWriter

from app.services.document_converter import convert_file_to_markdown

# 1. Test DOCX creation and conversion
doc = Document()
doc.add_heading("Docx Test Heading", level=1)
doc.add_paragraph("This is a test paragraph inside a Word document that was previously empty.")
docx_buf = io.BytesIO()
doc.save(docx_buf)
docx_bytes = docx_buf.getvalue()

docx_res = convert_file_to_markdown(docx_bytes, "test.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
print("DOCX Conversion:")
print("Suggested Title:", docx_res["suggested_title"])
print("Text snippet:", docx_res["text"][:100])
assert "Docx Test Heading" in docx_res["text"]
assert len(docx_res["text"]) > 10

# 2. Test PPTX creation and conversion
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Presentation Slide 1"
slide.placeholders[1].text = "First bullet point\nSecond bullet point"
pptx_buf = io.BytesIO()
prs.save(pptx_buf)
pptx_bytes = pptx_buf.getvalue()

pptx_res = convert_file_to_markdown(pptx_bytes, "presentation.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
print("\nPPTX Conversion:")
print("Suggested Title:", pptx_res["suggested_title"])
print("Text snippet:", pptx_res["text"][:120])
assert "Presentation Slide 1" in pptx_res["text"]
assert "First bullet point" in pptx_res["text"]

# 3. Test PDF sequential images using pypdf blank pages
writer = PdfWriter()
writer.add_blank_page(width=400, height=500)
writer.add_blank_page(width=400, height=500)
pdf_buf = io.BytesIO()
writer.write(pdf_buf)
pdf_bytes = pdf_buf.getvalue()

pdf_img_res = convert_file_to_markdown(pdf_bytes, "doc.pdf", "application/pdf", pages_as_images=True)
print("\nPDF Sequential Images Conversion:")
print("Suggested Title:", pdf_img_res["suggested_title"])
print("Number of page images:", len(pdf_img_res.get("page_images") or []))
print("Text snippet:", pdf_img_res["text"][:150])
assert len(pdf_img_res.get("page_images") or []) == 2
assert "![Page 1]" in pdf_img_res["text"]

print("\n>>> ALL BACKEND CONVERTER UNIT TESTS PASSED WITH 100% SUCCESS! <<<")
